import os
import subprocess
import tempfile
from pathlib import Path

from tools import WORKSPACE_ROOT


def _validate_path(path: str) -> str | None:
    full = os.path.join(WORKSPACE_ROOT, path) if not os.path.isabs(path) else path
    full = os.path.normpath(full)
    if not full.startswith(WORKSPACE_ROOT):
        return None
    if not os.path.isfile(full):
        return None
    return full


def read_pdf(path: str) -> str:
    full = _validate_path(path)
    if full is None:
        return "Error: file not found or access denied"
    try:
        from pypdf import PdfReader
        reader = PdfReader(full)
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except ImportError:
        pass
    try:
        from pdfminer.high_level import extract_text
        return extract_text(full)
    except ImportError:
        pass
    return "Error: no PDF library available (install pypdf or pdfminer.six)"


def read_docx(path: str) -> str:
    full = _validate_path(path)
    if full is None:
        return "Error: file not found or access denied"
    try:
        from docx import Document
        doc = Document(full)
        return "\n".join(p.text for p in doc.paragraphs)
    except ImportError:
        return "Error: python-docx library not available"
    except Exception as e:
        return f"Error reading docx: {e}"


def read_xlsx(path: str) -> str:
    full = _validate_path(path)
    if full is None:
        return "Error: file not found or access denied"
    try:
        import openpyxl
        wb = openpyxl.load_workbook(full, data_only=True)
        lines = []
        for sheet in wb.worksheets:
            lines.append(f"--- Sheet: {sheet.title} ---")
            for row in sheet.iter_rows(values_only=True):
                lines.append("\t".join(str(c) if c is not None else "" for c in row))
        return "\n".join(lines)
    except ImportError:
        return "Error: openpyxl library not available"
    except Exception as e:
        return f"Error reading xlsx: {e}"


def read_image_text(path: str) -> str:
    full = _validate_path(path)
    if full is None:
        return "Error: file not found or access denied"
    try:
        from PIL import Image
        import pytesseract
        return pytesseract.image_to_string(Image.open(full))
    except ImportError:
        return "Error: pytesseract or PIL not available"
    except Exception as e:
        return f"Error reading image: {e}"


EXT_READERS = {
    ".pdf": read_pdf,
    ".docx": read_docx,
    ".xlsx": read_xlsx,
    ".png": read_image_text,
    ".jpg": read_image_text,
    ".jpeg": read_image_text,
    ".tiff": read_image_text,
    ".bmp": read_image_text,
}


def read_document(path: str) -> str:
    ext = Path(path).suffix.lower()
    reader = EXT_READERS.get(ext)
    if reader is None:
        return f"Error: unsupported file extension '{ext}'"
    return reader(path)
